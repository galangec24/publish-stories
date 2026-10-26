import io
import json

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from rest_framework.test import APITestCase

from accounts.models import User
from payments.models import Subscription

from . import delta
from .importers import import_manuscript
from .models import Chapter, Column, Genre, Story


def doc(*lines):
    return delta.paragraphs_to_ops([([(line, None)], None) for line in lines])


class DeltaTests(TestCase):
    def test_validate_rejects_non_insert_ops(self):
        with self.assertRaises(ValueError):
            delta.validate_ops([{"delete": 3}])


class ImporterTests(TestCase):
    def test_markdown_front_matter_and_chapters(self):
        text = (
            "---\ntitle: Ang Diwata\nauthor: L. Cruz\ngenres: Fantasy, Philippine Folklore\n---\n"
            "Prologue text.\n\n# Unang Kabanata\n\nShe **ran** into the *forest*.\n\n## Scene\n\n- one\n- two\n\n---\n\n# Ikalawa\n\nEnd."
        )
        story = import_manuscript("diwata.md", text.encode())
        self.assertEqual(story.title, "Ang Diwata")
        self.assertEqual(story.author_name, "L. Cruz")
        self.assertEqual(story.genres, ["Fantasy", "Philippine Folklore"])
        self.assertEqual([c.title for c in story.chapters], ["Chapter 1", "Unang Kabanata", "Ikalawa"])
        ops = story.chapters[1].content
        self.assertIn({"insert": "ran", "attributes": {"bold": True}}, ops)
        self.assertIn({"insert": "forest", "attributes": {"italic": True}}, ops)
        self.assertIn({"insert": "\n", "attributes": {"list": "bullet"}}, ops)

    def test_text_chapter_detection(self):
        text = "Title: Gabi\n\nKabanata 1\nMadilim ang gabi.\nTahimik.\n\nKabanata 2\n\nUmaga na."
        story = import_manuscript("gabi.txt", text.encode())
        self.assertEqual(story.title, "Gabi")
        self.assertEqual([c.title for c in story.chapters], ["Kabanata 1", "Kabanata 2"])
        self.assertEqual(delta.plain_text(story.chapters[0].content), "Madilim ang gabi. Tahimik.\n")

    def test_docx(self):
        import docx

        document = docx.Document()
        document.add_paragraph("Word Story", style="Title")
        document.add_heading("Opening", level=1)
        paragraph = document.add_paragraph("Plain and ")
        paragraph.add_run("bold").bold = True
        document.add_heading("Second", level=1)
        document.add_paragraph("Item", style="List Bullet")
        buffer = io.BytesIO()
        document.save(buffer)

        story = import_manuscript("story.docx", buffer.getvalue())
        self.assertEqual(story.title, "Word Story")
        self.assertEqual([c.title for c in story.chapters], ["Opening", "Second"])
        self.assertIn({"insert": "bold", "attributes": {"bold": True}}, story.chapters[0].content)
        self.assertIn({"insert": "\n", "attributes": {"list": "bullet"}}, story.chapters[1].content)

    def test_filename_fallback_title(self):
        self.assertEqual(import_manuscript("my_first-story.txt", b"Hello").title, "My First Story")


class ReaderPaywallTests(APITestCase):
    def setUp(self):
        self.reader = User.objects.create_user(email="r@example.com", password="x", email_verified=True)
        self.genre = Genre.objects.create(name="Romance")
        self.story = Story.objects.create(title="Paid", price=49, free_chapters=1, status="published")
        self.story.genres.add(self.genre)
        Chapter.objects.create(story=self.story, title="One", order=0, content=doc("p1"))
        Chapter.objects.create(story=self.story, title="Two", order=1, content=doc("p2", "p3"))
        Chapter.objects.create(story=self.story, title="Three", order=2, content=doc("p4"))

    def test_free_chapters_are_whole_and_the_rest_withheld(self):
        data = self.client.get(f"/api/stories/{self.story.slug}/").data
        self.assertTrue(data["locked"])
        chapters = data["chapters"]
        self.assertEqual([c["title"] for c in chapters], ["One", "Two", "Three"])
        self.assertFalse(chapters[0]["locked"])
        self.assertEqual(delta.plain_text(chapters[0]["content"]), "p1\n")
        self.assertEqual([c["locked"] for c in chapters[1:]], [True, True])
        self.assertEqual([c["content"] for c in chapters[1:]], [None, None])

        self.story.free_chapters = 2
        self.story.save()
        chapters = self.client.get(f"/api/stories/{self.story.slug}/").data["chapters"]
        self.assertEqual(delta.plain_text(chapters[1]["content"]), "p2\np3\n")
        self.assertIsNone(chapters[2]["content"])

        self.story.free_chapters = 3  # every chapter free: nothing left to sell
        self.story.save()
        self.assertFalse(self.client.get(f"/api/stories/{self.story.slug}/").data["locked"])

    def test_paid_subscription_unlocks(self):
        Subscription.objects.create(user=self.reader, story=self.story, amount=49, status="paid")
        self.client.force_authenticate(self.reader)
        data = self.client.get(f"/api/stories/{self.story.slug}/").data
        self.assertFalse(data["locked"])
        self.assertEqual(delta.plain_text(data["chapters"][2]["content"]), "p4\n")

    def test_public_column_is_free(self):
        self.story.column = Column.objects.create(title="Editorial", is_public=True)
        self.story.save()
        self.assertFalse(self.client.get(f"/api/stories/{self.story.slug}/").data["locked"])

    def test_drafts_hidden_from_readers(self):
        draft = Story.objects.create(title="Secret", status="draft")
        self.assertEqual(self.client.get(f"/api/stories/{draft.slug}/").status_code, 404)

    def test_genre_filter_and_for_you(self):
        Story.objects.create(title="Other", status="published")
        listed = self.client.get("/api/stories/?genre=romance").data["results"]
        self.assertEqual([s["slug"] for s in listed], [self.story.slug])
        self.reader.preferred_genres.add(self.genre)
        self.client.force_authenticate(self.reader)
        home = self.client.get("/api/home/").data
        self.assertEqual([s["slug"] for s in home["for_you"]], [self.story.slug])
        self.assertEqual(home["by_genre"][0]["genre"]["slug"], "romance")


class AdminStudioTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_user(email="admin@example.com", password="x", is_admin=True)
        self.client.force_authenticate(self.admin)
        self.fantasy = Genre.objects.create(name="Fantasy")

    def test_non_admin_forbidden(self):
        reader = User.objects.create_user(email="r@example.com", password="x")
        self.client.force_authenticate(reader)
        self.assertEqual(self.client.get("/api/admin/stories/").status_code, 403)

    def test_create_story_edit_chapters_and_publish(self):
        response = self.client.post(
            "/api/admin/stories/",
            {"title": "Bagong Kwento", "genre_ids": [self.fantasy.id], "price": "49.00", "free_chapters": 2},
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        story_id = response.data["id"]
        chapter = response.data["chapters"][0]
        self.assertEqual(response.data["genres"][0]["slug"], "fantasy")

        self.assertEqual(self.client.post(f"/api/admin/stories/{story_id}/publish/").status_code, 400)

        saved = self.client.patch(
            f"/api/admin/chapters/{chapter['id']}/",
            {"content": doc("Once upon a time"), "revision": chapter["revision"]},
            format="json",
        )
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.data["revision"], 2)
        stale = self.client.patch(
            f"/api/admin/chapters/{chapter['id']}/", {"title": "x", "revision": 1}, format="json"
        )
        self.assertEqual(stale.status_code, 409)

        second = self.client.post(f"/api/admin/stories/{story_id}/chapters/", {"title": "Prologue", "order": 0}, format="json")
        self.assertEqual(second.status_code, 201)
        titles = [c["title"] for c in self.client.get(f"/api/admin/stories/{story_id}/chapters/").data]
        self.assertEqual(titles, ["Prologue", "Chapter 1"])

        reordered = self.client.post(
            f"/api/admin/stories/{story_id}/chapters/reorder/", {"order": [chapter["id"], second.data["id"]]}, format="json"
        )
        self.assertEqual([c["title"] for c in reordered.data], ["Chapter 1", "Prologue"])

        published = self.client.post(f"/api/admin/stories/{story_id}/publish/")
        self.assertEqual(published.data["status"], "published")
        self.assertEqual(published.data["word_count"], 4)

    def test_price_below_paymongo_minimum_rejected(self):
        response = self.client.post("/api/admin/stories/", {"title": "Cheap", "price": "5"}, format="json")
        self.assertEqual(response.status_code, 400)

    def test_bulk_import_with_genre_tagging(self):
        horror = Genre.objects.create(name="Horror")
        files = [
            SimpleUploadedFile("one.md", b"---\ntitle: Una\ngenres: Fantasy, Nonexistent\n---\n# Ch\n\nHello world."),
            SimpleUploadedFile("two.txt", b"Chapter 1\n\nText here."),
            SimpleUploadedFile("bad.pdf", b"%PDF"),
        ]
        manifest = [{"filename": "two.txt", "title": "Pangalawa", "genre_ids": [horror.id], "price": "25"}]
        response = self.client.post(
            "/api/admin/stories/import/",
            {"files": files, "genre_ids": str(self.fantasy.id), "manifest": json.dumps(manifest), "free_chapters": "1"},
            format="multipart",
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual((response.data["created"], response.data["failed"]), (2, 1))
        one, two, bad = response.data["results"]
        self.assertEqual(one["title"], "Una")
        self.assertEqual(one["genres"], ["Fantasy"])
        self.assertIn("Nonexistent", one["warnings"][0])
        self.assertEqual(two["title"], "Pangalawa")
        self.assertEqual(two["genres"], ["Horror"])
        self.assertFalse(bad["ok"])

        story = Story.objects.get(pk=two["story_id"])
        self.assertEqual(story.price, 25)
        self.assertEqual(story.free_chapters, 1)
        self.assertEqual(story.status, "draft")
        self.assertEqual(story.chapters.get().title, "Chapter 1")


class ViewCountTests(APITestCase):
    def setUp(self):
        from django.core.cache import cache

        cache.clear()
        self.story = Story.objects.create(title="Popular", status="published")
        self.reader = User.objects.create_user(email="r@example.com", password="x", email_verified=True)

    def view(self, **headers):
        return self.client.post(f"/api/stories/{self.story.slug}/view/", **headers)

    def test_one_view_per_reader_per_day_and_distinct_readers(self):
        guest = {"HTTP_X_VIEWER_ID": "a1b2c3d4-e5f6-4711-8899-aabbccddeeff"}
        self.assertEqual(self.view(**guest).data, {"view_count": 1, "reader_count": 1})
        self.assertEqual(self.view(**guest).data, {"view_count": 1, "reader_count": 1})  # same day, same reader

        self.client.force_authenticate(self.reader)
        self.assertEqual(self.view().data, {"view_count": 2, "reader_count": 2})

        # The next day the same reader adds a view but not a new reader.
        from .models import StoryView

        StoryView.objects.filter(user=self.reader).update(date="2000-01-01")
        self.assertEqual(self.view().data, {"view_count": 3, "reader_count": 2})

        detail = self.client.get(f"/api/stories/{self.story.slug}/").data
        self.assertEqual((detail["view_count"], detail["reader_count"]), (3, 2))
        stats = self.client.get("/api/stats/").data
        self.assertEqual(stats["total_readers"], 2)
        self.assertEqual(stats["total_views"], 3)

    def test_admin_previews_are_not_counted(self):
        admin = User.objects.create_user(email="a@example.com", password="x", is_admin=True)
        self.client.force_authenticate(admin)
        self.assertEqual(self.view().data, {"view_count": 0, "reader_count": 0})

    def test_drafts_cannot_be_viewed(self):
        draft = Story.objects.create(title="Hidden", status="draft")
        self.assertEqual(self.client.post(f"/api/stories/{draft.slug}/view/").status_code, 404)

    def test_reading_mode_preference(self):
        self.client.force_authenticate(self.reader)
        response = self.client.patch("/api/user/me/", {"reading_mode": "book"}, format="json")
        self.assertEqual(response.data["reading_mode"], "book")
        self.assertEqual(self.client.patch("/api/user/me/", {"reading_mode": "scroll3d"}, format="json").status_code, 400)


class AboutProfileTests(APITestCase):
    def test_public_read_has_defaults_and_admin_can_edit(self):
        data = self.client.get("/api/about/").data
        self.assertEqual(data["name"], "Maria Gemima Valderrama")
        self.assertEqual(len(data["roles"]), 4)
        self.assertIn("sunstar.com.ph/author/maria-gemima-c-valderrama", data["links"][0]["url"])

        admin = User.objects.create_user(email="a@example.com", password="x", is_admin=True)
        self.client.force_authenticate(admin)
        body = {**data, "tagline": "Writer · Davao", "roles": [{"icon": "pen", "title": "Author", "text": "Writes books."}]}
        body.pop("icons")
        saved = self.client.put("/api/admin/about/", body, format="json")
        self.assertEqual(saved.status_code, 200, saved.data)
        self.client.force_authenticate(None)
        public = self.client.get("/api/about/").data
        self.assertEqual(public["tagline"], "Writer · Davao")
        self.assertEqual(public["roles"][0]["title"], "Author")

    def test_validation_and_permissions(self):
        reader = User.objects.create_user(email="r@example.com", password="x")
        self.client.force_authenticate(reader)
        self.assertEqual(self.client.put("/api/admin/about/", {}, format="json").status_code, 403)
        admin = User.objects.create_user(email="a@example.com", password="x", is_admin=True)
        self.client.force_authenticate(admin)
        current = self.client.get("/api/admin/about/").data
        bad = {k: current[k] for k in ["name", "tagline", "intro", "closing", "photo_url", "roles", "links", "blog_heading", "blog_intro", "blog_cards"]}
        bad["roles"] = [{"icon": "rocket", "title": "x", "text": ""}]
        bad["links"] = [{"label": "x", "url": "javascript:alert(1)"}]
        response = self.client.put("/api/admin/about/", bad, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertIn("roles", response.data)
        self.assertIn("links", response.data)


class ChapterRulesTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_user(email="admin@example.com", password="x", is_admin=True)
        self.story = Story.objects.create(title="Rules", price=49, free_chapters=1, status="published")
        self.one = Chapter.objects.create(story=self.story, title="Chapter 1", order=0, content=doc("a"), word_count=1)
        self.two = Chapter.objects.create(story=self.story, title="Chapter 2", order=1, content=doc("b"), word_count=1)

    def test_chapter_titles_are_unique_within_a_story(self):
        self.client.force_authenticate(self.admin)
        url = f"/api/admin/stories/{self.story.id}/chapters/"
        self.assertIn("title", self.client.post(url, {"title": " chapter  1 "}, format="json").data)
        self.assertEqual(self.client.post(url, {"title": "Chapter 3"}, format="json").status_code, 201)
        renamed = self.client.patch(f"/api/admin/chapters/{self.two.id}/", {"title": "CHAPTER 1"}, format="json")
        self.assertEqual(renamed.status_code, 400)
        self.assertEqual(self.client.patch(f"/api/admin/chapters/{self.two.id}/", {"title": "Chapter 2"}, format="json").status_code, 200)
        other = Story.objects.create(title="Other")
        self.assertEqual(self.client.post(f"/api/admin/stories/{other.id}/chapters/", {"title": "Chapter 1"}, format="json").status_code, 201)

    def test_empty_chapters_after_the_free_ones_are_locked_too(self):
        Chapter.objects.create(story=self.story, title="Chapter 3", order=2, content=[], word_count=0)
        Chapter.objects.create(story=self.story, title="Chapter 4", order=3, content=[], word_count=0)
        self.story.free_chapters = 3
        self.story.save()
        chapters = self.client.get(f"/api/stories/{self.story.slug}/").data["chapters"]
        self.assertEqual([c["locked"] for c in chapters], [False, False, False, True])
        self.assertIsNone(chapters[3]["content"])

    def test_import_makes_repeated_headings_distinct(self):
        from .serializers import unique_titles

        self.assertEqual(unique_titles(["Prologue", "One", "prologue", "Prologue"]), ["Prologue", "One", "prologue (2)", "Prologue (3)"])



class RecentlyDeletedChaptersTests(APITestCase):
    def test_delete_is_restorable_in_place(self):
        from datetime import timedelta

        from django.utils import timezone

        admin = User.objects.create_user(email="admin@example.com", password="x", is_admin=True)
        self.client.force_authenticate(admin)
        story = Story.objects.create(title="Trash", status="published")
        one, two, three = (Chapter.objects.create(story=story, title=f"Chapter {i}", order=i - 1, content=doc(f"p{i}")) for i in (1, 2, 3))

        self.assertEqual(self.client.delete(f"/api/admin/chapters/{two.id}/").status_code, 204)
        self.assertEqual([c["title"] for c in self.client.get(f"/api/stories/{story.slug}/").data["chapters"]], ["Chapter 1", "Chapter 3"])
        trash = self.client.get(f"/api/admin/stories/{story.id}/chapters/deleted/").data
        self.assertEqual([(t["title"], t["excerpt"]) for t in trash], [("Chapter 2", "p2")])

        # A new chapter takes the old name meanwhile; the restored one gets "(restored)".
        self.client.post(f"/api/admin/stories/{story.id}/chapters/", {"title": "Chapter 2"}, format="json")
        restored = self.client.post(f"/api/admin/chapters/{two.id}/restore/").data
        self.assertEqual(restored["title"], "Chapter 2 (restored)")
        titles = [c["title"] for c in self.client.get(f"/api/stories/{story.slug}/").data["chapters"]]
        self.assertEqual(titles, ["Chapter 1", "Chapter 2 (restored)", "Chapter 3", "Chapter 2"])
        self.assertEqual(self.client.get(f"/api/admin/stories/{story.id}/chapters/deleted/").data, [])
        self.assertEqual(self.client.post(f"/api/admin/chapters/{two.id}/restore/").status_code, 404)  # not deleted

        # After 30 days the trash is emptied.
        self.client.delete(f"/api/admin/chapters/{three.id}/")
        Chapter.all_objects.filter(pk=three.pk).update(deleted_at=timezone.now() - timedelta(days=31))
        self.assertEqual(self.client.get(f"/api/admin/stories/{story.id}/chapters/deleted/").data, [])
        self.assertFalse(Chapter.all_objects.filter(pk=three.pk).exists())



class ChapterHeadingsTests(APITestCase):
    def test_headings_are_kept_on_save_and_listed_for_every_chapter(self):
        story = Story.objects.create(title="Sections")
        ops = [
            {"insert": "Part One"}, {"insert": "\n", "attributes": {"header": 1}},
            {"insert": "Body text\nThe Storm"}, {"insert": "\n", "attributes": {"header": 2}},
            {"insert": "  "}, {"insert": "\n", "attributes": {"header": 3}},  # empty heading: skipped
        ]
        chapter = Chapter.objects.create(story=story, title="One", order=0, content=ops)
        Chapter.objects.create(story=story, title="Two", order=1, content=doc("plain"))
        self.assertEqual(chapter.headings, [{"level": 1, "text": "Part One"}, {"level": 2, "text": "The Storm"}])
        admin = User.objects.create_user(email="admin@example.com", password="x", is_admin=True)
        self.client.force_authenticate(admin)
        listed = self.client.get(f"/api/admin/stories/{story.id}/").data["chapters"]
        self.assertEqual([c["headings"] for c in listed], [chapter.headings, []])
        saved = self.client.patch(f"/api/admin/chapters/{chapter.id}/", {"content": [{"insert": "New"}, {"insert": "\n", "attributes": {"header": 2}}]}, format="json").data
        self.assertEqual(saved["headings"], [{"level": 2, "text": "New"}])


class BoldTitleTests(TestCase):
    def test_short_fully_bold_lines_count_as_sections(self):
        ops = [
            {"insert": "The Woman Who Came Back", "attributes": {"bold": True}}, {"insert": "\n"},
            {"insert": "She walked in. "}, {"insert": "Loudly", "attributes": {"bold": True}}, {"insert": ".\n"},  # partly bold: no
            {"insert": " ".join(["word"] * 20), "attributes": {"bold": True}}, {"insert": "\n"},  # too long: no
            {"insert": "Item", "attributes": {"bold": True}}, {"insert": "\n", "attributes": {"list": "bullet"}},  # list: no
            {"insert": "Part Two"}, {"insert": "\n", "attributes": {"header": 1}},
        ]
        self.assertEqual(delta.headings(ops), [{"level": 2, "text": "The Woman Who Came Back"}, {"level": 1, "text": "Part Two"}])
