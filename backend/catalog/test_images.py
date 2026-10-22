import io

from django.core.files.uploadedfile import SimpleUploadedFile
from PIL import Image
from rest_framework.test import APITestCase

from accounts.models import User

from . import delta
from .models import Chapter, Story


def png_bytes(w=2400, h=1200):
    buffer = io.BytesIO()
    Image.new("RGB", (w, h), (200, 80, 40)).save(buffer, "PNG")
    return buffer.getvalue()


class StoryImageTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_user(email="a@example.com", password="x", is_admin=True)
        self.client.force_authenticate(self.admin)

    def test_upload_resizes_and_serves_webp(self):
        response = self.client.post(
            "/api/admin/images/", {"image": SimpleUploadedFile("pic.png", png_bytes(), "image/png")}, format="multipart"
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual((response.data["width"], response.data["height"]), (1600, 800))
        path = response.data["url"].split("testserver")[1]
        served = self.client.get(path)
        self.assertEqual(served.status_code, 200)
        self.assertEqual(served["Content-Type"], "image/webp")
        self.assertIn("immutable", served["Cache-Control"])

    def test_rejects_non_images(self):
        response = self.client.post(
            "/api/admin/images/", {"image": SimpleUploadedFile("x.png", b"not an image", "image/png")}, format="multipart"
        )
        self.assertEqual(response.status_code, 400)

    def test_chapters_accept_uploaded_images_but_not_pasted_data(self):
        story = Story.objects.create(title="Pics")
        chapter = Chapter.objects.create(story=story, title="One")
        ok = [
            {"insert": "Before the picture.\n"},
            {"insert": {"image": "https://api.example.com/api/media/images/abc.webp"}},
            {"insert": "\n", "attributes": {"align": "center"}},
            {"insert": "After the picture.\n"},
        ]
        saved = self.client.patch(f"/api/admin/chapters/{chapter.pk}/", {"content": ok}, format="json")
        self.assertEqual(saved.status_code, 200, saved.data)
        for bad in ({"image": "data:image/png;base64,AAAA"}, {"image": "javascript:alert(1)"}, {"video": "https://x"}):
            response = self.client.patch(
                f"/api/admin/chapters/{chapter.pk}/", {"content": [{"insert": bad}, {"insert": "\n"}]}, format="json"
            )
            self.assertEqual(response.status_code, 400, bad)

    def test_uploads_are_admin_only(self):
        reader = User.objects.create_user(email="r@example.com", password="x")
        self.client.force_authenticate(reader)
        response = self.client.post(
            "/api/admin/images/", {"image": SimpleUploadedFile("pic.png", png_bytes(10, 10), "image/png")}, format="multipart"
        )
        self.assertEqual(response.status_code, 403)


def docx_with_picture():
    import docx

    document = docx.Document()
    document.add_paragraph("My Book", style="Title")
    document.add_heading("Kabanata 1", level=1)
    document.add_paragraph("The river was quiet.")
    picture = io.BytesIO()
    Image.new("RGB", (300, 200), (30, 120, 200)).save(picture, "PNG")
    picture.seek(0)
    document.add_picture(picture)
    document.add_paragraph("Then the rain came.")
    document.add_page_break()
    document.add_paragraph("A new page begins.")
    document.add_heading("Kabanata 2", level=1)
    document.add_paragraph("Morning.")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


class ImportOneTests(APITestCase):
    def setUp(self):
        from catalog.models import Genre

        self.admin = User.objects.create_user(email="a@example.com", password="x", is_admin=True)
        self.client.force_authenticate(self.admin)
        self.genre = Genre.objects.create(name="Drama")

    def post(self, **fields):
        upload = SimpleUploadedFile("book.docx", docx_with_picture())
        return self.client.post("/api/admin/stories/import-one/", {"file": upload, **fields}, format="multipart")

    def test_preview_saves_nothing_and_reports_chapters_and_images(self):
        from catalog.models import StoryImage

        data = self.post(preview="true").data
        self.assertEqual(data["title"], "My Book")
        self.assertEqual([c["title"] for c in data["chapters"]], ["Kabanata 1", "Kabanata 2"])
        self.assertEqual(data["image_count"], 1)
        self.assertEqual(Story.objects.count(), 0)
        self.assertEqual(StoryImage.objects.count(), 0)

    def test_split_modes(self):
        by_page = [c["title"] for c in self.post(preview="true", split="page_breaks").data["chapters"]]
        self.assertEqual(len(by_page), 2)  # one page break
        none = self.post(preview="true", split="none").data["chapters"]
        self.assertEqual(len(none), 1)

    def test_create_keeps_pictures_titles_and_fonts(self):
        import json

        from catalog.models import StoryImage

        response = self.post(
            split="auto",
            chapter_titles=json.dumps(["The River", ""]),
            story=json.dumps({"title": "Renamed", "genre_ids": [self.genre.id], "font_family": "lora", "font_size": 20, "price": "49"}),
        )
        self.assertEqual(response.status_code, 201, response.data)
        story = Story.objects.get()
        self.assertEqual((story.title, story.font_family, story.font_size), ("Renamed", "lora", 20))
        self.assertEqual([c.title for c in story.chapters.all()], ["The River", "Kabanata 2"])
        self.assertEqual(StoryImage.objects.count(), 1)
        first = story.chapters.first().content
        images = [op for op in first if isinstance(op["insert"], dict)]
        self.assertEqual(len(images), 1)
        self.assertIn("/api/media/images/", images[0]["insert"]["image"])
        # The picture sits between the two paragraphs it was between in Word.
        text = delta.plain_text(first)
        self.assertLess(text.index("quiet"), text.index("rain"))

    def test_invalid_font_size_rolls_back_pictures(self):
        import json

        from catalog.models import StoryImage

        response = self.post(story=json.dumps({"font_size": 99}))
        self.assertEqual(response.status_code, 400)
        self.assertEqual((Story.objects.count(), StoryImage.objects.count()), (0, 0))
